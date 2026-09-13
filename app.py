from flask import Flask, render_template, request, redirect, url_for, flash
import sqlite3
import os
from werkzeug.utils import secure_filename
import openpyxl

app = Flask(__name__)
app.secret_key = "sixsense-coders-secret-key"

UPLOAD_FOLDER = "static/uploads"
EXCEL_FILE = "users.xlsx"
COLLEGE_EXCEL = "college_admins.xlsx"
COMPANY_EXCEL = "companies.xlsx"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# --- Excel Functions ---
def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(["ID", "Name", "Email", "Password", "Resume"])
        wb.save(EXCEL_FILE)

def init_college_excel():
    if not os.path.exists(COLLEGE_EXCEL):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CollegeAdmins"
        ws.append(["ID", "College Name", "Admin Name", "Email", "Password"])
        wb.save(COLLEGE_EXCEL)

def init_company_excel():
    if not os.path.exists(COMPANY_EXCEL):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Companies"
        ws.append(["ID", "Company Name", "HR Name", "Email", "Password"])
        wb.save(COMPANY_EXCEL)

def save_to_excel(name, email, password, resume_filename):
    init_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[2].value == email:
            return False
    new_id = ws.max_row
    ws.append([new_id, name, email, password, resume_filename])
    wb.save(EXCEL_FILE)
    return True

def save_college_to_excel(college_name, admin_name, email, password):
    init_college_excel()
    wb = openpyxl.load_workbook(COLLEGE_EXCEL)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[3].value == email:
            return False
    new_id = ws.max_row
    ws.append([new_id, college_name, admin_name, email, password])
    wb.save(COLLEGE_EXCEL)
    return True

def save_company_to_excel(company_name, hr_name, email, password):
    init_company_excel()
    wb = openpyxl.load_workbook(COMPANY_EXCEL)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[3].value == email:
            return False
    new_id = ws.max_row
    ws.append([new_id, company_name, hr_name, email, password])
    wb.save(COMPANY_EXCEL)
    return True

def check_login_excel(email, password):
    if not os.path.exists(EXCEL_FILE):
        return None
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] == email and row[3] == password:
            return {"id": row[0], "name": row[1], "email": row[2], "resume": row[4]}
    return None

def check_college_login_excel(email, password):
    # Pehle hardcoded admin check, fir excel check
    if email == "admin@college.edu" and password == "admin123":
        return {"college": "Demo College", "name": "Super Admin", "email": email}
    if not os.path.exists(COLLEGE_EXCEL):
        return None
    wb = openpyxl.load_workbook(COLLEGE_EXCEL)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3] == email and row[4] == password:
            return {"id": row[0], "college": row[1], "name": row[2], "email": row[3]}
    return None

def check_company_login_excel(email, password):
    if email == "company@sixsense.com" and password == "company123":
        return {"company": "SixSense Demo", "name": "HR", "email": email}
    if not os.path.exists(COMPANY_EXCEL):
        return None
    wb = openpyxl.load_workbook(COMPANY_EXCEL)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3] == email and row[4] == password:
            return {"id": row[0], "company": row[1], "name": row[2], "email": row[3]}
    return None

# --- SQLite ---
def get_db():
    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    return conn

def create_table():
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            resume TEXT
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS college_admins (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            college_name TEXT NOT NULL,
            admin_name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company_name TEXT NOT NULL,
            hr_name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL
        )
    """)
    conn.commit()
    conn.close()
    init_excel()
    init_college_excel()
    init_company_excel()

# --- Routes ---
@app.route('/')
def home():
    return render_template('myweb.html')

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form["name"]
        email = request.form["email"]
        password = request.form["password"]
        resume = request.files.get("resume")
        resume_filename = None
        if resume and resume.filename!= "":
            resume_filename = secure_filename(resume.filename)
            resume.save(os.path.join(app.config["UPLOAD_FOLDER"], resume_filename))
        saved = save_to_excel(name, email, password, resume_filename)
        if not saved:
            return "Email already registered! <a href='/login'>Login karo</a>"
        try:
            conn = get_db()
            conn.execute("INSERT INTO students (name, email, password, resume) VALUES (?,?,?,?)",
                         (name, email, password, resume_filename))
            conn.commit()
            conn.close()
        except:
            pass
        return redirect(url_for('login'))
    return render_template("register.html")

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == "POST":
        email = request.form["email"]
        password = request.form["password"]
        user = check_login_excel(email, password)
        if user:
            return render_template('dashboard.html', user=user)
        else:
            return "Invalid email or password! <a href='/login'>Try again</a>"
    return render_template("login.html")

@app.route('/dashboard')
def dashboard():
    return render_template('dashboard.html', user={"name": "User"})

# =========== COLLEGE ADMIN ===========
@app.route('/college-admin/register', methods=['GET', 'POST'])
def college_admin_register():
    if request.method == "POST":
        college_name = request.form.get("college_name")
        admin_name = request.form.get("admin_name")
        email = request.form.get("email")
        password = request.form.get("password")

        saved = save_college_to_excel(college_name, admin_name, email, password)
        if not saved:
            return "College Email already registered! <a href='/college-admin/login'>Login karo</a>"

        try:
            conn = get_db()
            conn.execute("INSERT INTO college_admins (college_name, admin_name, email, password) VALUES (?,?,?,?)",
                         (college_name, admin_name, email, password))
            conn.commit()
            conn.close()
        except:
            pass
        return redirect(url_for('college_admin_login'))
    return render_template('college_admin_register.html')

@app.route('/college-admin/login', methods=['GET', 'POST'])
def college_admin_login():
    if request.method == "POST":
        email = request.form.get("email")
        password = request.form.get("password")
        admin = check_college_login_excel(email, password)
        if admin:
            return redirect(url_for('college_admin_dashboard'))
        else:
            return "Invalid College Admin! <a href='/college-admin/login'>Try again</a>"
    return render_template('college_admin_login.html')

@app.route('/college-admin/dashboard')
def college_admin_dashboard():
    conn = get_db()
    students = conn.execute("SELECT * FROM students").fetchall()
    colleges = conn.execute("SELECT * FROM college_admins").fetchall()
    conn.close()
    return render_template('college_admin_dashboard.html', students=students, colleges=colleges, total_students=len(students), total_colleges=len(colleges))

# =========== COMPANY PANEL ===========
@app.route('/company/register', methods=['GET', 'POST'])
def company_register():
    if request.method == "POST":
        company_name = request.form.get("company_name")
        hr_name = request.form.get("hr_name")
        email = request.form.get("email")
        password = request.form.get("password")

        saved = save_company_to_excel(company_name, hr_name, email, password)
        if not saved:
            return "Company Email already registered! <a href='/company/login'>Login karo</a>"
        try:
            conn = get_db()
            conn.execute("INSERT INTO companies (company_name, hr_name, email, password) VALUES (?,?,?,?)",
                         (company_name, hr_name, email, password))
            conn.commit()
            conn.close()
        except:
            pass
        return redirect(url_for('company_login'))
    return render_template('company_register.html')

@app.route('/company/login', methods=['GET', 'POST'])
def company_login():
    if request.method == "POST":
        email = request.form.get("email")
        password = request.form.get("password")
        comp = check_company_login_excel(email, password)
        if comp:
            return redirect(url_for('company_dashboard'))
        else:
            return "Invalid Company Login! <a href='/company/login'>Try again</a>"
    return render_template('company_login.html')

@app.route('/company/dashboard')
def company_dashboard():
    conn = get_db()
    students = conn.execute("SELECT * FROM students").fetchall()
    conn.close()
    return render_template('company_dashboard.html', students=students)

if __name__ == "__main__":
    create_table()
    app.run(debug=True)